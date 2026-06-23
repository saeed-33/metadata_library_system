#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate an Arabic Word document explaining the implementation of
entities, persistence models, DbContexts, and the borrow operation
in the Metadata Library System.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_rtl(paragraph):
    """Set paragraph direction to RTL."""
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_heading_rtl(doc, text, level=1):
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:asciiTheme'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsiTheme'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:cstheme'), 'Arial')
    run.font.size = Pt(16 if level == 1 else (14 if level == 2 else 12))
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    set_rtl(heading)
    return heading


def add_paragraph_rtl(doc, text, bold=False, italic=False, size=12, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:asciiTheme'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsiTheme'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:cstheme'), 'Arial')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    set_rtl(p)
    return p


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run._element.rPr.rFonts.set(qn('w:asciiTheme'), 'Courier New')
    run._element.rPr.rFonts.set(qn('w:hAnsiTheme'), 'Courier New')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0, 100, 0)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(6)
    # Keep LTR for code
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_bullet_rtl(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:asciiTheme'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsiTheme'), 'Arial')
    run.font.size = Pt(12)
    set_rtl(p)
    return p


def main():
    doc = Document()

    # Title
    title = doc.add_heading(level=0)
    run = title.add_run('شرح طريقة تنفيذ الكيانات، النماذج، السياقات، وعملية الاستعارة')
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:asciiTheme'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsiTheme'), 'Arial')
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    set_rtl(title)

    add_paragraph_rtl(doc, 'مشروع Metadata Library System – نظام إدارة مكتبة', bold=True, size=13)
    add_paragraph_rtl(doc, 'تم إعداد هذا الملف لشرح كيفية تنفيذ الطبقات الأساسية في المشروع مع الإشارة إلى الأخطاء المنطقية المكتشفة.')
    doc.add_paragraph()

    # 1. Project architecture
    add_heading_rtl(doc, '1. نظرة عامة على بنية المشروع', level=1)
    add_paragraph_rtl(doc,
        'يقسم المشروع إلى طبقات متعددة وفق أسلوب Clean Architecture المستوحى من CQRS. الطبقات الرئيسية هي:')
    add_bullet_rtl(doc, 'LibrarySystem.Domain: تحتوي على الكيانات (Entities) والتعدادات (Enums) والعقود الأساسية.')
    add_bullet_rtl(doc, 'LibrarySystem.Application: تحتوي على أوامر واستعلامات CQRS، والتحقق، وملفات التعيين (Mapping Profiles).')
    add_bullet_rtl(doc, 'LibrarySystem.DataAccess: تحتوي على نماذج قاعدة البيانات (Persistence Models)، وسياقات Entity Framework، والمستودعات.')
    add_bullet_rtl(doc, 'LibrarySystem.API: الطبقة الخارجية التي تستضيف المتحكمات (Controllers) وإعدادات DI والـ Middleware.')
    doc.add_paragraph()

    # 2. Domain Entities
    add_heading_rtl(doc, '2. تنفيذ الكيانات في طبقة Domain', level=1)
    add_paragraph_rtl(doc,
        'تقع الكيانات داخل المجلد LibrarySystem.Domain/entities/. جميع الكيانات ترث من صنف أساسي اسمه BaseEntity.')

    add_heading_rtl(doc, '2.1 الصنف BaseEntity', level=2)
    add_paragraph_rtl(doc, 'يحتوي BaseEntity على الخصائص المشتركة لكل الكيانات:')
    add_code_block(doc, '''public abstract class BaseEntity : ISoftDelete
{
    public int Id { get; set; }
    public DateTime CreatedAt { get; set; }
    public string? CreatedBy { get; set; }
    public DateTime? ModifiedAt { get; set; }
    public string? ModifiedBy { get; set; }
    public bool IsDeleted { get; set; }
    public DateTime? DeletedAt { get; set; }
}''')
    add_paragraph_rtl(doc,
        'الكيان يطبق واجهة ISoftDelete التي تحتوي على IsDeleted وDeletedAt، مما يتيح الحذف المنطقي (Soft Delete). ملاحظة: BaseEntity لا يطبق IAuditable رغم وجود حقول التدقيق، وهذا يسبب مشكلة في تعبئة الحقول آلياً كما سنرى لاحقاً.')

    add_heading_rtl(doc, '2.2 التسلسل الهرمي للموارد Resource', level=2)
    add_paragraph_rtl(doc,
        'يوجد صنف مجرد اسمه Resource يرث من BaseEntity، ويمثل أي مورد في النظام (مادة، وسائط، مجموعة مواد). الكيانات التي ترث منه هي:')
    add_bullet_rtl(doc, 'Item: المادة الأساسية في المكتبة.')
    add_bullet_rtl(doc, 'Media: الملفات والصور المرتبطة بالمادة.')
    add_bullet_rtl(doc, 'ItemSet: مجموعة من المواضيع المتعلقة ببعضها.')
    add_code_block(doc, '''public abstract class Resource : BaseEntity
{
    public string Type { get; protected set; } = null!;
    public int? OwnerId { get; set; }
    public virtual SystemUser? Owner { get; set; }
    private readonly List<Value> _values = new();
    public virtual IReadOnlyCollection<Value> Values => _values.AsReadOnly();
}''')

    add_heading_rtl(doc, '2.3 كيانات الدورة Circulation', level=2)
    add_paragraph_rtl(doc, 'الكيانات الأساسية لعملية الإعارة هي:')
    add_code_block(doc, '''public class BorrowRecord : BaseEntity
{
    public int CopyId { get; set; }
    public ItemCopy Copy { get; set; } = null!;
    public int PatronId { get; set; }
    public Patron Patron { get; set; } = null!;
    public DateTime BorrowDate { get; set; }
    public DateTime DueDate { get; set; }
    public DateTime? ReturnDate { get; set; }
    public string Status { get; set; } = "Active";
}''')
    add_code_block(doc, '''public class ItemCopy : BaseEntity
{
    public int ItemId { get; set; }
    public Item Item { get; set; } = null!;
    public string Barcode { get; set; } = string.Empty;
    public int Status { get; set; }   // 0: متاح, 1: معار, 2: مرجعي فقط, 3: ؟
    public string? Notes { get; set; }
}''')
    add_paragraph_rtl(doc,
        'الملاحظة هنا أن Status في ItemCopy يستخدم أرقاماً سحرية (Magic Numbers) بدلاً من enum، مما يجعل الكود هشاً وعرضة للأخطاء.')
    doc.add_paragraph()

    # 3. Persistence Models
    add_heading_rtl(doc, '3. تنفيذ نماذج قاعدة البيانات Persistence Models', level=1)
    add_paragraph_rtl(doc,
        'تقع النماذج داخل LibrarySystem.DataAccess/Persistence/models/. المشروع يستخدم نموذجين متوازيين: كيانات Domain ونماذج Persistence. الهدف هو فصل المنطق عن التخزين، لكنه يضيف تعقيداً كبيراً.')

    add_heading_rtl(doc, '3.1 BasePersistenceModel', level=2)
    add_code_block(doc, '''public abstract class BasePersistenceModel : ISoftDelete
{
    [Key] public int Id { get; set; }
    [Required] public DateTime CreatedAt { get; set; }
    [StringLength(128)] public string? CreatedBy { get; set; }
    public DateTime? ModifiedAt { get; set; }
    [StringLength(128)] public string? ModifiedBy { get; set; }
    public bool IsDeleted { get; set; }
    public DateTime? DeletedAt { get; set; }
}''')
    add_paragraph_rtl(doc,
        'مثل BaseEntity، لا يطبق BasePersistenceModel واجهة IAuditable، فلا يتم تعبئة حقول التدقيق تلقائياً إلا إذا عُبأت يدوياً.')

    add_heading_rtl(doc, '3.2 استراتيجية TPT للوراثة', level=2)
    add_paragraph_rtl(doc,
        'يستخدم المشروع Table-per-Type (TPT) للتعامل مع الوراثة. يوجد جدول Resources الأساسي، وجداول منفصلة Items, Medias, ItemSets ترث منه.')
    add_code_block(doc, '''[Table("Resources")]
public abstract class ResourceModel : BasePersistenceModel
{
    [Required, StringLength(50)] public string Type { get; set; } = string.Empty;
    public int? OwnerId { get; set; }
    [ForeignKey("OwnerId")] public SystemUserModel? Owner { get; set; }
    public ICollection<ValueModel> Values { get; set; } = new List<ValueModel>();
}

[Table("Items")]
public class ItemModel : ResourceModel
{
    public int? TemplateId { get; set; }
    public ResourceTemplateModel? Template { get; set; }
    public virtual ICollection<ItemCopyModel> Copies { get; set; } = new HashSet<ItemCopyModel>();
    public virtual ICollection<MediaModel> Medias { get; set; } = new List<MediaModel>();
    public virtual ICollection<ItemSetModel> ItemSets { get; set; } = new List<ItemSetModel>();
}''')

    add_heading_rtl(doc, '3.3 ملاحظات على النماذج', level=2)
    add_bullet_rtl(doc, 'يوجد أخطاء إملائية في أسماء الملفات: MdeiaModel.cs و BookmarkModel .cs (فراغ في النهاية).')
    add_bullet_rtl(doc, 'لا يوجد RoleModel رغم وجود كيان Role في Domain.')
    add_bullet_rtl(doc, 'BorrowRecord.cs و Patron.cs في مجلد models يحملان أسماء بدون لاحقة Model، مما يسبب لبساً.')
    doc.add_paragraph()

    # 4. DbContexts
    add_heading_rtl(doc, '4. تنفيذ سياقات قاعدة البيانات DbContexts', level=1)

    add_heading_rtl(doc, '4.1 ApplicationDbContext', level=2)
    add_paragraph_rtl(doc,
        'يستضيف جميع نماذج التطبيق الرئيسية. يقوم بثلاث مهام أساسية:')
    add_bullet_rtl(doc, 'تطبيق فلتر عالمي Global Query Filter لإخفاء الصفوف المحذوفة منطقياً.')
    add_bullet_rtl(doc, 'تحويل عمليات Delete إلى Soft Delete داخل SaveChangesAsync.')
    add_bullet_rtl(doc, 'تدقيق الوقت CreatedAt/ModifiedAt للكيانات التي تطبق IAuditable.')
    add_code_block(doc, '''public override async Task<int> SaveChangesAsync(CancellationToken ct = default)
{
    foreach (var entry in ChangeTracker.Entries<ISoftDelete>())
    {
        if (entry.State == EntityState.Deleted)
        {
            entry.State = EntityState.Modified;
            entry.Entity.IsDeleted = true;
            entry.Entity.DeletedAt = DateTime.UtcNow;
        }
    }
    foreach (var entry in ChangeTracker.Entries<IAuditable>())
    {
        if (entry.State == EntityState.Added) entry.Entity.CreatedAt = DateTime.UtcNow;
        else if (entry.State == EntityState.Modified) entry.Entity.ModifiedAt = DateTime.UtcNow;
    }
    return await base.SaveChangesAsync(ct);
}''')
    add_paragraph_rtl(doc,
        'المشكلة: لأن BasePersistenceModel لا يطبق IAuditable، فإن معظم الكيانات تحصل على DateTime.MinValue في CreatedAt ما لم يتم تعبئتها يدوياً.')

    add_heading_rtl(doc, '4.2 CustomIdentityDbContext', level=2)
    add_paragraph_rtl(doc,
        'يستخدم لتخزين المستخدمين والأدوار عبر ASP.NET Core Identity. يحتوي على AppUserModel و AppRoleModel. لكنه يتجاهل (Ignore) نماذج التطبيق مثل SystemUserModel رغم أن AppUserModel.SystemUserId يفترض أن يشير إليها.')
    add_code_block(doc, '''protected override void OnModelCreating(ModelBuilder builder)
{
    base.OnModelCreating(builder);
    builder.Entity<AppRoleModel>().HasQueryFilter(role => !role.IsDeleted);
    builder.Ignore<SystemUserModel>();
    builder.Ignore<ResourceModel>();
    builder.Ignore<ItemModel>();
    // ...
}''')
    add_paragraph_rtl(doc,
        'هذا يعني أن العلاقة بين مستخدم الهوية والمستخدم في تطبيق المكتبة هي علاقة عبر قاعدتين منفصلتين دون قيد foreign key على مستوى SQL.')

    add_heading_rtl(doc, '4.3 LoggingDbContext', level=2)
    add_paragraph_rtl(doc,
        'سياق بسيط جداً في الميكروسيرفيس Logging.API، يحتوي على جدول Logs فقط.')
    doc.add_paragraph()

    # 5. Mapping
    add_heading_rtl(doc, '5. التعيين بين Domain و Persistence', level=1)
    add_paragraph_rtl(doc,
        'يستخدم المشروع AutoMapper للتعيين بين الكيانات ونماذجها. توجد ملفات تعيين في Application/Mappings/ (للـ DTOs) و DataAccess/Mappings/ (للـ Persistence).')
    add_code_block(doc, '''CreateMap<Item, ItemModel>()
    .ForMember(dest => dest.ItemSets, opt => opt.Ignore())
    .ForMember(dest => dest.Copies, opt => opt.MapFrom(src => src.Copies))
    .ReverseMap()
    .ForMember(dest => dest.ItemSets, opt => opt.Ignore())
    .ForMember(dest => dest.Values, opt => opt.Ignore())
    .ForMember(dest => dest.Copies, opt => opt.Ignore());

CreateMap<BorrowRecord, BorrowRecordModel>().ReverseMap().MaxDepth(2);''')
    add_paragraph_rtl(doc,
        'المستودع MappedRepository<TDomain, TPersistence> يلف DbContext ويقوم بترجمة التعبيرات من Domain إلى Persistence باستخدام ExpressionTranslator. كما يحفظ قائمة بالإضافات المعلقة pendingAdds لمزامنة المفاتيح المتولدة تلقائياً مع الكائنات في طبقة Domain.')
    doc.add_paragraph()

    # 6. Borrow operation
    add_heading_rtl(doc, '6. تنفيذ عملية الاستعارة Borrow Operation', level=1)

    add_heading_rtl(doc, '6.1 Checkout (إعارة نسخة)', level=2)
    add_paragraph_rtl(doc,
        'الأمر CheckoutCommand يستقبل Barcode و PatronId وتاريخ استحقاق اختياري. مساره: LibrarySystem.Application/Commands/Circulation/CheckoutCommandHandler.cs')
    add_paragraph_rtl(doc, 'خطوات المعالج:')
    add_bullet_rtl(doc, 'يبحث عن النسخة باستخدام Barcode (لكنه يحمل كل النسخ إلى الذاكرة!).')
    add_bullet_rtl(doc, 'يفحص أن Status == 0 (متاح).')
    add_bullet_rtl(doc, 'يتحقق من وجود Template وأن IsBorrowable == true.')
    add_bullet_rtl(doc, 'يحسب تاريخ الاستحقاق من Template.DefaultBorrowDays أو إعداد GlobalBorrowDays.')
    add_bullet_rtl(doc, 'ينشئ BorrowRecord بحالة Active.')
    add_bullet_rtl(doc, 'يغير حالة النسخة إلى 1 (معار).')
    add_bullet_rtl(doc, 'يحفظ التغييرات.')
    add_code_block(doc, '''var allCopies = await _unitOfWork.ItemCopies.GetAllAsync();
var copy = allCopies.FirstOrDefault(c => c.Barcode == request.Barcode);

if (copy.Status != 0) throw new InvalidOperationException("Copy is not available.");

var item = await _unitOfWork.Items.GetByIdAsync(copy.ItemId);
if (!item.TemplateId.HasValue) throw ...;
var template = await _unitOfWork.ResourceTemplates.GetByIdAsync(item.TemplateId.Value);
if (template == null || !template.IsBorrowable) throw ...;

var borrowRecord = _mapper.Map<BorrowRecord>(request);
borrowRecord.CopyId = copy.Id;
borrowRecord.BorrowDate = DateTime.UtcNow;
borrowRecord.DueDate = dueDate;
borrowRecord.Status = "Active";

copy.Status = 1;
await _unitOfWork.BorrowRecords.AddAsync(borrowRecord);
_unitOfWork.ItemCopies.Update(copy);
await _unitOfWork.SaveChangesAsync();''')

    add_heading_rtl(doc, '6.2 Return (إرجاع نسخة)', level=2)
    add_paragraph_rtl(doc,
        'الأمر ReturnCommand يستقبل Barcode فقط. مساره: LibrarySystem.Application/Commands/Circulation/ReturnCommand.cs')
    add_bullet_rtl(doc, 'يحمل كل النسخ والسجلات إلى الذاكرة.')
    add_bullet_rtl(doc, 'يبحث عن السجل النشط المطابق للنسخة وغير المُرجع.')
    add_bullet_rtl(doc, 'يضع ReturnDate ويغير Status إلى Returned.')
    add_bullet_rtl(doc, 'يستعيد حالة النسخة إلى 0 دون مراعاة حالتها السابقة.')
    add_code_block(doc, '''var allCopies = await _unitOfWork.ItemCopies.GetAllAsync();
var copy = allCopies.FirstOrDefault(c => c.Barcode == request.Barcode);

var allRecords = await _unitOfWork.BorrowRecords.GetAllAsync();
var activeRecord = allRecords.FirstOrDefault(r => r.CopyId == copy.Id && r.ReturnDate == null);

activeRecord.ReturnDate = DateTime.UtcNow;
activeRecord.Status = "Returned";
copy.Status = 0;

_unitOfWork.BorrowRecords.Update(activeRecord);
_unitOfWork.ItemCopies.Update(copy);
await _unitOfWork.SaveChangesAsync();''')

    add_heading_rtl(doc, '6.3 الاستعلامات المرتبطة', level=2)
    add_bullet_rtl(doc, 'GetActiveLoansQuery: يسترجع السجلات غير المُرجعة مع تضمين Copy.Item و Patron.')
    add_bullet_rtl(doc, 'GetBorrowHistoryQuery: يسترجع كل السجلات.')
    add_bullet_rtl(doc, 'GetOverdueLoansQuery: يحسب السجلات المتأخرة في الذاكرة دون تحديث Status.')
    doc.add_paragraph()

    # 7. Errors and logic issues
    add_heading_rtl(doc, '7. الأخطاء والمشاكل المنطقية المكتشفة', level=1)

    add_heading_rtl(doc, '7.1 أخطاء حرجة Critical', level=2)
    add_bullet_rtl(doc,
        'Docker Compose يستخدم متغير ConnectionStrings__DefaultConnection للـ Logging.API بينما الكود يقرأ ConnectionStrings:Default. هذا يعني فشل تشغيل Logging.API في Docker.')
    add_bullet_rtl(doc,
        'نقص التحكم بالتزامن في Checkout: يُحمل كل الجدول إلى الذاكرة ويفحص Status. طلبان متزامنان لنفس الباركود يمكنهما إنشاء سجلي إعارة نشطين لنفس النسخة.')
    add_bullet_rtl(doc,
        'التحكمات لا تفرق بين الأدوار: أي مستخدم مصدق يمكنه الإعارة والإرجاع والإنشاء والحذف، وليس فقط المسؤول/أمين المكتبة.')
    add_bullet_rtl(doc,
        'Logging.API لا يتطلب مصادقة، لذا يمكن لأي عميل إرسال أو قراءة السجلات.')

    add_heading_rtl(doc, '7.2 أخطاء منطقية في الاستعارة', level=2)
    add_bullet_rtl(doc,
        'Checkout لا يتحقق من وجود PatronId؛ يمكن إنشاء إعارة لمشترك غير موجود.')
    add_bullet_rtl(doc,
        'لا يوجد حد أقصى لعدد الإعارات لكل مشترك رغم وجود ActiveLoansCount.')
    add_bullet_rtl(doc,
        'Return لا يتحقق من Status == "Active"، بل يكتفي بـ ReturnDate == null.')
    add_bullet_rtl(doc,
        'Return يعيد النسخة دائماً إلى Status=0 حتى لو كانت مرجعية أو تحت الصيانة.')
    add_bullet_rtl(doc,
        'لا يوجد enum للحالات؛ يُستخدم int سحري في ItemCopy و string سحري في BorrowRecord.')
    add_bullet_rtl(doc,
        'حالة Overdue لا تُحدّث أبداً في قاعدة البيانات، بل تُحسب وقت الاستعلام فقط.')
    add_bullet_rtl(doc,
        'ItemTitle في DTOs يُعرض كـ ItemId.ToString() وليس كعنوان فعلي.')

    add_heading_rtl(doc, '7.3 مشاكل عامة في الأداء والبيانات', level=2)
    add_bullet_rtl(doc,
        'العديد من المعالجات تستخدم GetAllAsync() ثم تُفلتر في الذاكرة مما يسبب بطء شديد عند كبر البيانات.')
    add_bullet_rtl(doc,
        'لا توجد معاملات Transaction عند التعديلات المتعددة (CreateItem, CreateMedia, Register).')
    add_bullet_rtl(doc,
        'الحذف المنطقي لا ينتقل للبيانات التابعة مثل Values و Copies و Medias.')
    add_bullet_rtl(doc,
        'حقول CreatedBy/ModifiedBy لا تُعبأ أبداً.')
    add_bullet_rtl(doc,
        'MappedRepository.Update ينشئ نموذج Persistence جديداً ويستدعي Update، مما قد يمحو خصائص لم تُحمّل.')
    add_bullet_rtl(doc,
        'RequestLoggingMiddleware ينتظر HTTP كاملاً إلى Logging.API قبل إنهاء الطلب، مما يسبب تأخراً.')

    add_heading_rtl(doc, '7.4 مشاكل بنيوية وتسميات', level=2)
    add_bullet_rtl(doc, 'أسماء ملفات تحتوي على فراغات أو أخطاء إملائية: MdeiaModel.cs, BookmarkModel .cs, UpdateMediaComman.cs, UnitOfWork .cs, IUnitOfWork .cs.')
    add_bullet_rtl(doc, 'اختلاف namespace: Domain.common/entities مقابل Domain.Common/Entities.')
    add_bullet_rtl(doc, 'اضطراب في إصدارات .NET: المشاريع تستهدف net8.0 بينما Dockerfiles و CI تستخدم .NET 9.')
    add_bullet_rtl(doc, 'حزم NuGet تحتوي على ثغرات معروفة: AutoMapper 13.0.1 و SixLabors.ImageSharp 3.1.5.')
    doc.add_paragraph()

    # 8. Summary
    add_heading_rtl(doc, '8. الخلاصة', level=1)
    add_paragraph_rtl(doc,
        'المشروع يتبع بنية نظيفة ومنظمة، لكن هناك عدة مشاكل منطقية وعملية تحتاج إصلاحاً عاجلاً. أهمها:')
    add_bullet_rtl(doc, 'تصحيح التحقق من الوجود والتزامن في عملية الإعارة.')
    add_bullet_rtl(doc, 'استبدال الأرقام/النصوص السحرية بـ enums.')
    add_bullet_rtl(doc, 'إضافة المعاملات Transactions والتحكم بالتزامن.')
    add_bullet_rtl(doc, 'تقييد الصلاحيات حسب الأدوار.')
    add_bullet_rtl(doc, 'إصلاح متغير Docker Compose وإعدادات Logging.API.')
    add_bullet_rtl(doc, 'تصحيح أسماء الملفات وتوحيد namespaces وإصدارات .NET.')

    # Save
    output_path = 'Arabic_Implementation_Explanation.docx'
    doc.save(output_path)
    print(f'Saved: {output_path}')


if __name__ == '__main__':
    main()
